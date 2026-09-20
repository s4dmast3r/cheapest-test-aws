"""Informe auditado: métricas originales, incertidumbre y evidencia trazable."""
from pathlib import Path
import csv, json, re, statistics
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

LAB=Path('lab2'); OUT=LAB/'report'; d=Document()
audit=json.loads((OUT/'audit-results.json').read_text())
assert audit['complete_runs']==44
config={r['run']:r for r in json.loads((OUT/'results-summary.json').read_text())}
rows=[dict(config[r['run']],**r) for r in audit['runs']]
ref=[]
for f in (LAB/'refinement').glob('*-r*.json'):
    if f.name.endswith(('.command.json','.config.json')):continue
    r=json.loads(f.read_text())
    if 'overall' in r:
        m=re.fullmatch(r'(get|post)-u(\d+)-r(\d+)',f.stem)
        ref.append(dict(run=f.stem,endpoint=m[1].upper(),users=int(m[2]),rep=int(m[3]),**r['overall']))
ref.sort(key=lambda r:(r['endpoint'],r['users'],r['rep']))
def fmt(n,dp=2):return f'{n:.{dp}f}'
def rng(rs,key,dp=2):return f"{fmt(min(r[key] for r in rs),dp)}–{fmt(max(r[key] for r in rs),dp)}"
def select(method,scenario):return [r for r in rows if r['endpoint']==method and r['scenario']==scenario]
styles=d.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Calibri';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(11)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.08
styles['Title'].font.size=Pt(27)
styles['Heading 1'].font.size=Pt(19)
styles['Heading 2'].font.size=Pt(13)
styles['Heading 3'].font.size=Pt(11)
styles['Caption'].font.size=Pt(9)
for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    prop=styles[name].element.find(qn('w:pPr'))
    if prop is not None:
        for border in list(prop.findall(qn('w:pBdr'))):prop.remove(border)
    styles[name].font.underline=False
def setup(sec,land=False):
    sec.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
    sec.page_width=Inches(11 if land else 8.5);sec.page_height=Inches(8.5 if land else 11)
    sec.top_margin=sec.bottom_margin=Inches(.6 if land else .7)
    sec.left_margin=sec.right_margin=Inches(.6 if land else .75)
    sec.header_distance=sec.footer_distance=Inches(.3)
setup(d.sections[0])
foot=d.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=foot.add_run('Cheapest · Laboratorio 2   |   ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(text,style=None):return d.add_paragraph(text,style)
def h(text,level=1):return d.add_heading(text,level)
def page(title=None,land=None):
    if land is None:d.add_page_break()
    else:setup(d.add_section(WD_SECTION.NEW_PAGE),land)
    if title:h(title)
def table(headers,data,widths=None,size=10):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    available=9.8 if d.sections[-1].orientation==WD_ORIENT.LANDSCAPE else 7
    widths=widths or [available/len(headers)]*len(headers)
    widths=[x*available/sum(widths) for x in widths]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for ri,values in enumerate([headers]+list(data)):
        row=t.rows[0] if ri==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr();split=OxmlElement('w:cantSplit');trpr.append(split)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge,val in [('top',55),('bottom',55),('left',75),('right',75)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');marg.append(el)
            cp.append(marg)
            if ri==0 or ri%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8F0' if ri==0 else 'F6F8FA');cp.append(sh)
            para=cell.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
            r=para.add_run(str(value));r.font.size=Pt(size);r.bold=ri==0
    p('') .paragraph_format.space_after=Pt(0)
    return t
figure=0
def img(name,title,explanation,width=7):
    global figure
    path=LAB/name;assert path.exists(),path
    para=d.add_paragraph();para.paragraph_format.keep_with_next=True
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline=para.add_run().add_picture(str(path),width=Inches(width))
    inline._inline.docPr.set('descr',title+'. '+explanation)
    figure+=1
    p(f'Figura {figure}. {title}. {explanation}','Caption')

h('Pruebas de carga del monolito Cheapest',0)
p('Laboratorio 2 · Arquitectura de Software · Universidad de los Andes','Subtitle')
p('Rama lab2-jose · Campaña principal del 7 de septiembre de 2026 · Revisión con pruebas adicionales')
h('1 Introducción y objetivos')
p('Se evalúan la consulta de productos disponibles y la creación de pedidos del monolito NestJS con PostgreSQL. El objetivo es medir el cumplimiento de los dos ASR, reconocer la degradación bajo carga y proponer mejoras que respondan a la evidencia del código y de las ejecuciones.')
p('La matriz principal contiene 44 corridas completas: 22 GET y 22 POST. El GET cumple su límite de latencia en las cinco corridas normales de JMeter, pero incumple desde el primer nivel Python de la matriz. El POST conserva 0 % de errores hasta 7500 usuarios y supera 74 % en el pico de 18000. El rendimiento útil cae aunque el contador total de solicitudes terminadas aumente.')
p('Los usuarios nominales de JMeter no equivalen a los usuarios que repiten solicitudes en Python. Por ello, se añaden pruebas con el mismo ramp-up y duración para estudiar el inicio del incumplimiento; los resultados cercanos al límite se presentan con su variabilidad, sin atribuir una capacidad productiva exacta al experimento local.')
h('2 ASR evaluados')
table(['ASR','Operación y contexto','Criterio'],[
    ['ASR 1','GET de productos pedidos históricamente, en promoción y disponibles en la zona. Operación normal: 500 req/min.','p99 < 1000 ms'],
    ['ASR 2','POST de creación de pedidos durante promociones. Referencia: 5000 req/min.','Error ≤ 2 %']], [1,4.5,1.5])
p('Rutas: GET /logistics/tenderos/productos-disponibles, con tiendaId y zona; POST /logistics/pedidos. Se considera éxito HTTP 200 y HTTP 201, respectivamente. En el pico, un timeout no prueba que el pedido no haya sido persistido: la entrega distingue confirmación al cliente y creación observada en la base.')
p('La guía, el warm-up y el checklist de correcciones se conservan en reference/. Las secciones 6 y 7 presentan mediciones e interpretación; la 8 revisa los cuellos de botella; la 9 responde las cuatro preguntas; las secciones 12 y 13 contienen prompts y evidencia.')

page('3 Entorno y reproducibilidad')
table(['Componente','Configuración registrada'],[
 ['Equipo','Apple M1 · 8 CPU · 16 GB RAM · macOS 26.6.2'],
 ['Aplicación','NestJS local · Node 26.8.1 · compilación de producción'],
 ['Base','PostgreSQL 15 Alpine en Docker · 8 CPU disponibles'],
 ['Generadores','Apache JMeter 5.6.3 con Java 21; Python con aiohttp'],
 ['Conexiones','Proveedor TypeORM sin límite explícito; pg-pool usa 10'],
 ['Repositorio','cheapest-api · rama lab2-jose'],
 ['Estado inicial','Commit 32aeb398c41c1a38b0e606b3cb4e6e628c2f846b'],
 ['Salud','GET /health → HTTP 200 y Health Check OK']], [1.45,5.55])
p('La aplicación, la base y el generador comparten el computador. El contenedor de prueba es cheapest-postgres; la campaña principal usa la base cheapest y el puerto 3000. Las pruebas adicionales restauran una base separada, cheapest_lab2_refinement, y usan el puerto 3002. No se cambiaron los endpoints para mejorar su desempeño durante la medición.')
h('Preparación y controles',2)
p('Se inspeccionaron load-seed.yaml y database-seeder.service.ts antes de definir las dos variantes. Se compiló el backend, se verificó /health y se conservaron los snapshots GET y POST. JMeter se ejecutó en modo no gráfico; la descarga de la distribución oficial fue verificada mediante SHA-512. El plan local conserva los parámetros y listeners exigidos por la guía.')
p('Antes de cada corrida POST principal se restauró el mismo snapshot de 4500 pedidos y 31500 ítems. Los archivos de conteos anteriores y posteriores permiten comprobarlo. GET no escribe; en las cargas mayores se reinició la aplicación para evitar trabajo pendiente de la corrida anterior. Las pruebas adicionales restauran el snapshot correspondiente antes de cada repetición.')
h('Cómo repetir',2)
p('El ZIP incluye semillas, planes, cuerpos, script Python, dependencias, configuraciones por corrida y snapshots. README_ENTREGA.md explica la organización y el procedimiento. Cada CSV original se conserva junto con su configuración y resumen; los consolidados results_get.csv y results_post.csv permiten volver a calcular las métricas sin depender del Word.')
p('Las versiones del entorno, los identificadores de prueba y la fecha de las promociones afectan la reproducción. Para comparar tratamientos se debe restaurar el mismo snapshot y conservar los mismos parámetros; al regenerar semillas pueden cambiar UUID y fechas. La ausencia de un calentamiento prolongado y el uso de recursos compartidos limitan la precisión del umbral.')

page('4 Diseño y distribución de datos')
p('Se modela una red urbana de 75 tiendas en cinco zonas, con 1800 productos. Son supuestos experimentales, no una estimación extraída de clientes reales. Ambas variantes evitan repartir toda la actividad por igual y mantienen suficiente volumen para que la consulta histórica y la persistencia tengan trabajo relevante.')
seed_data=[('Tiendas',75,75),('Zonas',5,5),('Productos',1800,1800),('Catálogos',125,125),('Catálogo producto',22500,22500),('Promociones',3600,5400),('Pedidos iniciales',18000,4500),('Ítems de pedido',126000,31500),('Despachos',13500,3375),('Disponibilidad por zona',22500,22500),('Notas crédito',540,135),('Ítems de inventario',4500,4500),('Registros de compra',9000,9000),('Registros de venta',18000,18000),('Productos externos',240,240),('Ventas',6000,6000),('Ítems de venta',18000,18000)]
table(['Cantidad objetivo','GET ASR 1','POST ASR 2'],seed_data,[3.5,1.75,1.75],10)
p('GET usa seeds/get-asr1.yaml: peso de cabeza 0.8 y fracción 0.2. POST usa seeds/post-asr2.yaml: peso 0.9 y fracción 0.2. Los conteos de tablas y las distribuciones materializadas se conservaron en los archivos *-baseline-*.csv; las relaciones promocion_tiendas incluyen además una asociación de la semilla base.')

page('4 Diseño y distribución de datos',None)
h('Distribución efectiva y propósito',2)
p('Para GET, el historial de 18000 pedidos y 126000 ítems obliga a resolver la condición de compra previa sobre una relación extensa. El sesgo 80/20 representa tiendas con mayor actividad y zonas densas. El plan SQL guardado consulta una tienda de cabeza de Zona Norte: selecciona 959 pedidos de 18000 y recorre 126000 ítems antes de devolver cuatro productos.')
p('Para POST, el historial inicial baja a 4500 pedidos, las promociones suben a 5400 y la concentración pasa a 90/20. El cuerpo Python contiene 24 productos diferentes, de modo que la validación por ítem sea significativa. El endpoint no consulta promociones ni actualiza stock durante la creación; por ello, aumentar promociones da contexto comercial al dataset, pero no agrega directamente consultas al camino del POST.')
table(['Propiedad observada','GET','POST'],[
 ['Catálogos en Zona Norte','100 de 125','112 de 125'],
 ['Catálogos en otras cuatro zonas','7, 6, 6, 6','4, 3, 3, 3'],
 ['Promociones activas','2520 de 3600','3780 de 5400'],
 ['Asociaciones promoción tienda','3601','5401'],
 ['Ítems por pedido del cliente Python','No aplica','24'],
 ['Ítems por pedido del cliente JMeter','No aplica','4']], [3.5,1.75,1.75])
h('Qué representa realmente Pareto',2)
p('El seeder utiliza una asignación determinista con diez posiciones: ocho o nueve favorecen la cabeza del 20 %. Ese mecanismo afecta las asignaciones de pedidos a tiendas, catálogos a tiendas y zonas, promociones a tiendas e ítems a productos. Otras relaciones siguen ciclos: promoción a producto y disponibilidad a catálogo/producto. El inventario conserva referencias de las tiendas iniciales. No se afirma que todas las columnas sigan una distribución Pareto continua.')
p('Además, las solicitudes de carga utilizan una tienda y zona fijas; el tráfico no recorre una población de tiendas con muestreo Pareto. Se mide un caso caliente repetido sobre una base heterogénea. Las fechas de pedidos avanzan un día por registro, por lo que el histórico GET cubre décadas; esto no cambia el filtro “alguna vez pedido”, pero sí limita el realismo temporal. El mock de tiendas siempre acepta la tienda y espera 10 ms.')
p('Una comparación uniforme frente a sesgada no fue ejecutada. El diseño permite explorar las hipótesis de lectura histórica y validación secuencial; no demuestra por sí solo que el sesgo cause la degradación. La Pregunta 2 explica los posibles sesgos y los controles necesarios.')

page('5 Diseño experimental')
table(['Escenario','Usuarios','Ramp-up s','Meseta s','Reps por endpoint'],[
 ['Smoke',5,5,'1 loop',1],['Baja',30,10,'1 loop',1],['Media',100,20,'1 loop',1],['Normal',450,50,'1 loop',5],['Alta',1500,75,60,3],['Muy alta',3000,100,60,3],['Estrés',7500,150,60,3],['Pico',18000,200,60,5]], [2.1,1,1.2,1.2,1.5])
p('JMeter ejecuta un envío por hilo en los primeros cuatro niveles. A 450 hilos y 50 s, las solicitudes se distribuyen aproximadamente a nueve por segundo; los CSV registran como máximo un hilo activo simultáneo. Por tanto, “450 hilos” no significa 450 solicitudes en vuelo. Python inicia usuarios gradualmente y cada uno repite sin pausa hasta terminar la meseta: es un modelo cerrado de concurrencia.')
h('Métricas y denominadores',2)
p('La tabla usa la corrida completa. Throughput total = solicitudes finalizadas / tiempo total observado; incluye fallos. Throughput exitoso cuenta solamente HTTP 200 GET o 201 POST. Error % cuenta respuestas inesperadas, errores HTTP y excepciones de cliente. El código 0 identifica ausencia de respuesta HTTP, no un estado del servidor.')
p('JMeter calcula percentiles con la interpolación compatible con su reporte HTML; Python usa rango más cercano, posición ceil(p × N). Se recalcularon los percentiles desde los CSV. La duración Python de 60 s corresponde a la meseta; el tiempo total incluye ramp-up y drenaje de solicitudes. El timeout configurado es 30 s, pero la demora del event loop puede hacer que la latencia observada exceda ese valor.')
p('Los JSON también contienen métricas de solicitudes iniciadas durante la meseta. Esa cohorte puede terminar después; su conteo dividido por 60 no equivale exactamente a las respuestas completadas dentro de esos 60 s. Por ello, la comparación principal utiliza el tiempo total y conserva los datos de fase para análisis posterior.')
h('Pruebas para reducir el intervalo',2)
p('Se usa el mismo cliente Python, ramp-up de 30 s, meseta de 60 s, cuerpo POST de 24 ítems y snapshot restaurado. Los niveles se exploran de forma adaptativa y se repiten cerca del límite. Esta segunda serie se presenta aparte: no se mezclan sus resultados con las 44 corridas de la matriz original ni se ignoran repeticiones discordantes.')

names={'smoke':'Smoke','baja':'Baja','media':'Media','normal':'Normal','alta':'Alta','muy-alta':'Muy alta','estres':'Estrés','pico':'Pico'}
for method in ['GET','POST']:
    page('6 Resultados '+method,True)
    p(f'Tabla {1 if method=="GET" else 2}. Las 22 corridas de {method}. Duración: meseta Python; “loop” indica una iteración JMeter. RPS incluye fallos. Separador decimal: punto.')
    chosen=[r for r in rows if r['endpoint']==method]
    values=[[names[r['scenario']],r['rep'],r['users'],r['ramp_up'],r['duration'] or 'loop',r['requests'],fmt(r['p95_ms']),fmt(r['p99_ms']),fmt(r['throughput_rps']),fmt(r['error_pct'])] for r in chosen]
    table(['Escenario','Rep','Usuarios','Ramp s','Dur s','Solicitudes','p95 ms','p99 ms','RPS total','Error %'],values,[1.25,.4,.75,.65,.6,.95,1.1,1.1,.95,.85],9)
    p('Fuente: raw/'+method.lower()+'-<escenario>-<repetición>.csv. Auditoría: report/audit-results.json. Las configuraciones por corrida contienen los parámetros completos.','Caption')

page('6 Resultados de las pruebas adicionales',True)
p('Tabla 3. Modelo Python común: ramp-up 30 s, meseta 60 s, snapshot restaurado por corrida. Estos resultados son adicionales a la matriz mínima.')
table(['Endpoint','Usuarios','Rep','Solicitudes','p95 ms','p99 ms','RPS total','Error %'],[[r['endpoint'],r['users'],r['rep'],r['requests'],fmt(r['p95_ms']),fmt(r['p99_ms']),fmt(r['throughput_rps']),fmt(r['error_pct'])] for r in ref],[1,.8,.5,1.2,1.3,1.3,1.1,.9],9)
p('La variación entre repeticiones forma parte del resultado. El límite se interpreta en la sección 7; un incumplimiento aislado no se convierte automáticamente en una capacidad máxima reproducible. Fuentes: refinement/*.csv y *.json.','Caption')

page('6 Gráficas de cumplimiento',False)
img('graphs/get-asr.png','p99 del GET en las 22 corridas','La línea de 1000 ms permite ver el incumplimiento desde alta carga Python. La separación vertical distingue los dos modelos de generación.',7)
img('graphs/post-asr.png','Errores del POST en las 22 corridas','Los puntos en cero están incluidos. Las cinco corridas de pico superan ampliamente la referencia del 2 %.',7)

page('6 Throughput y refinamiento')
img('graphs/throughput.png','Throughput total y exitoso','En el pico, terminar solicitudes con error aumenta el total, pero no la capacidad útil. Las curvas de JMeter y Python se presentan en paneles separados.',7)
img('graphs/refinement.png','Pruebas adicionales con un modelo común','Cada punto representa una corrida, incluidas las repeticiones discordantes. El eje GET es logarítmico para mostrar las variaciones de latencia.',7)

page('7 Análisis de resultados')
h('Punto de inflexión y primer incumplimiento',2)
p('En la matriz original, ASR 1 se incumple primero: el GET pasa de p99 de 44.49–46 ms en las cinco corridas normales a 5994–9609 ms con 1500 usuarios Python. Sin embargo, el cambio de un envío por hilo a usuarios que repiten impide deducir un intervalo de capacidad entre 450 y 1500. Las tasas objetivo de la guía, 8.33 req/s GET y 83.33 req/s POST, tampoco son equivalentes a esos usuarios nominales.')
g100=[r for r in ref if r['endpoint']=='GET' and r['users']==100]
g150=[r for r in ref if r['endpoint']=='GET' and r['users']==150]
if g100 and g150:
    bad=sum(r['p99_ms']>=1000 for r in g150)
    p(f'Con el modelo Python común, 100 usuarios cumple en {len(g100)} de {len(g100)} corridas, con p99 de {rng(g100,"p99_ms")} ms. A 150 usuarios hay {bad} incumplimientos en {len(g150)} corridas: p99 de {rng(g150,"p99_ms")} ms. El inicio observado de incumplimientos intermitentes queda entre 100 y 150 usuarios; 150 no es un límite estable, porque también registra una corrida que cumple.')
g300=[r for r in ref if r['endpoint']=='GET' and r['users']==300]
if g300:p(f'A 300 usuarios, {sum(r["p99_ms"]>=1000 for r in g300)} de {len(g300)} corridas supera 1000 ms, con p99 de {rng(g300,"p99_ms")} ms. Estos niveles delimitan una zona de transición local; no son un intervalo de confianza estadístico ni una garantía de capacidad en producción.')
postgroups={n:[r for r in ref if r['endpoint']=='POST' and r['users']==n] for n in sorted({r['users'] for r in ref if r['endpoint']=='POST'})}
if postgroups:
    p('En el refinamiento POST, los rangos de error por nivel son: '+ '; '.join(f'{n} usuarios, {rng(rs,"error_pct")} % ({len(rs)} corridas)' for n,rs in postgroups.items())+'.')
    good=[n for n,rs in postgroups.items() if len(rs)>=3 and all(r['error_pct']<=2 for r in rs)]
    bad=[n for n,rs in postgroups.items() if len(rs)>=3 and all(r['error_pct']>2 for r in rs)]
    if good and bad and max(good)<min(bad):p(f'Los extremos repetidos sitúan el cruce observado del ASR 2 entre {max(good)} y {min(bad)} usuarios en esta configuración. El extremo inferior cumple todas sus repeticiones y el superior las incumple todas. No se extrapola ese intervalo a otro ramp-up, generador o equipo.')
    else:p('Los niveles disponibles no delimitan todavía un cruce estable mediante dos extremos repetidos. Se conserva el intervalo amplio de la matriz principal y se reporta el refinamiento como evidencia de sensibilidad al ramp-up, sin inventar un umbral exacto.')
h('Comparación entre endpoints',2)
p('El GET degrada primero respecto a su criterio: mantiene respuestas HTTP exitosas mientras su cola de latencia ya viola el segundo. POST puede responder lentamente y todavía cumplir su ASR de errores; sus fallos aparecen más tarde, cuando predominan timeouts y problemas de conexión. Esta diferencia responde también a criterios distintos, no prueba que el POST sea una operación más eficiente.')

page('7 Degradación y alcance de las mediciones')
h('Saturación gradual y caída del rendimiento útil',2)
p('GET aumenta su p99 a 12.64–13.53 s con 3000 usuarios, y a 29.72–30.24 s con 7500. El throughput exitoso mediano permanece alrededor de 246–259 req/s entre alta carga y estrés: más concurrencia agrega espera sin crecimiento proporcional de trabajo útil. En el pico, la mediana cae a aproximadamente 101 éxitos/s y los errores llegan a 73.69–78.02 %.')
p('POST conserva cero errores hasta 7500 usuarios en la matriz, pero su p99 aumenta de 4.96–5.18 s a 26.16–28.22 s. El throughput exitoso mediano desciende de unos 299 a 282 req/s. A 18000 usuarios hay 74.24–77.23 % de errores y aproximadamente 101 éxitos/s. La latencia muestra degradación previa; el porcentaje de fallos cambia abruptamente al atravesar los límites temporales y de conexión.')
h('Qué significa el fallo de un pedido',2)
p('En toda la matriz GET se registran 425000 errores de cliente por timeout. POST registra 435262 errores: 416138 timeouts y 19124 errores de conexión. No se observaron códigos HTTP 4xx o 5xx en esas fallas; el cliente registró estado 0. Los ejemplos y conteos están en los CSV y en la auditoría.')
p('En POST pico 5 hubo 116097 solicitudes: 29523 respuestas 201 y 86574 errores. Después de detener el trabajo pendiente, la base tenía 41456 pedidos adicionales y 994944 ítems adicionales, exactamente 24 por pedido. La diferencia entre pedidos persistidos y respuestas confirmadas muestra resultados ambiguos para el cliente. No puede usarse la tasa HTTP como conteo exacto de ventas perdidas ni reintentar sin idempotencia.')
h('Límites de atribución',2)
p('El retraso p99 del event loop del generador alcanza 3865.80 ms en POST pico 5. Esto indica que el cliente local participa en la degradación y afecta la precisión de timeouts. La telemetría disponible no mide directamente la cola de adquisición del pool. Tampoco se ejecutó una comparación con índices nuevos, pool modificado o un generador remoto; por eso no se asigna un porcentaje causal a cada recurso.')
p('Las pruebas GET y POST se ejecutaron por separado y con distinto volumen inicial. Los POST de JMeter tienen cuatro ítems y los de Python 24. No se midieron tráfico mixto, caché fría controlada, sesiones de tiendas distintas ni una tasa abierta exacta de 5000 req/min. El experimento demuestra el comportamiento de esta configuración local, no certifica por sí solo los ASR en el entorno comercial completo.')

page('8 Análisis del cuello de botella')
h('GET y trabajo dentro de PostgreSQL',2)
p('Confirmado por código: producto.repository.ts envía una consulta con tres EXISTS para historial, promociones activas y disponibilidad por zona. No hace una consulta de aplicación por cada producto. La auditoría de índices confirma la ausencia de índices útiles en pedidos.tiendaId, items_pedido.productoId, catalogos.zona y disponibilidad_zona.productoId.')
p('Observado: EXPLAIN ANALYZE BUFFERS muestra hash joins y agregaciones. Recorre 126000 ítems, filtra 18000 pedidos y examina 22500 filas de disponibilidad; devuelve cuatro productos. El plan tarda 33.033 ms en ejecución, con 3238 bloques encontrados en caché. La rama del historial llega aproximadamente a 23 ms y es el mayor componente de ese plan aislado.')
p('Interpretación: el trabajo relacional es el candidato dominante del GET, apoyado por el plan y por la meseta de throughput. No se afirma que los EXISTS sean la causa única: PostgreSQL transforma la expresión y usa varias estrategias. Falta una comparación controlada con índices para demostrar cuánto de la latencia bajo concurrencia corresponde a ejecución SQL frente a espera.')
h('POST y validación secuencial',2)
p('Confirmado por código: PedidoService.create espera findById dentro del ciclo de ítems. Un pedido de 24 ítems realiza 24 lecturas secuenciales de validación antes de guardar, además de la espera del mock de tienda. PedidoRepository.create usa save con entidades en cascada. Esa llamada no permite deducir una inserción independiente por ítem: el ORM puede agrupar escrituras.')
p('Confirmado por configuración: database.providers.ts no fija el tamaño del pool; la versión instalada de pg-pool aplica diez conexiones. Hipótesis: la combinación de múltiples viajes y espera por conexiones amplifica la latencia. Observado: el POST acumula latencias y pierde confirmaciones a alta carga. No se instrumentó el tiempo de adquisición de conexión, por lo que no se declara probado que el pool sea el único cuello dominante.')
h('Evidencia que falta para decidir una causa única',2)
p('La CPU compartida y el retraso del generador impiden aislar servidor y cliente en el pico. Para GET existe evidencia directa del costo SQL; para POST existe evidencia del patrón secuencial, pero faltan trazas por consulta y por adquisición de conexión. La Pregunta 3 define los experimentos que distinguirían esas hipótesis.')

analysis=(OUT/'analysis-draft.md').read_text()
qparts=re.split(r'^## ',analysis,flags=re.M)
for i in range(1,5):
    section=next(x for x in qparts if x.startswith(f'Pregunta {i} '))
    lines=section.strip().splitlines();page('9 Respuestas a las preguntas 1 a 4' if i==1 else f'9 Pregunta {i}')
    h(lines[0],2)
    for line in lines[1:]:
        if line.strip():
            line=line.replace('Valida la hipótesis','Permite evaluar la hipótesis')
            p(line)

page('10 Propuestas arquitectónicas')
h('Primero mejorar los caminos actuales',2)
p('Antes de una promoción, priorizar la validación de productos por lote en POST, manteniendo una transacción breve para pedido e ítems. Incorporar una clave de idempotencia persistida con unicidad y recuperación del resultado: el cliente podrá repetir una solicitud ambigua sin duplicar el pedido. Medir el cambio con el mismo snapshot, cuerpo y carga antes de desplegarlo.')
p('Para GET, probar índices compuestos según selectividad y plan: pedidos(tiendaId, id), items_pedido(productoId, pedidoId), catalogos(zona, id) y disponibilidad_zona(productoId, catalogoId). Comparar también el orden alternativo pedidoId/productoId cuando el plan empiece por la tienda. El costo de mantener índices en las escrituras debe medirse; una lista de candidatos no sustituye EXPLAIN ni una prueba A/B.')
h('Control de concurrencia y recursos',2)
p('Instrumentar espera y ocupación del pool, latencia SQL y retraso del event loop. Aplicar límites de concurrencia y tiempos máximos coherentes entre cliente, API y base. Dimensionar las conexiones con un presupuesto global: replicar la API multiplica los pools. Un aumento ciego de diez a cientos de conexiones puede empeorar la contención cuando PostgreSQL ya está saturado.')
h('Evolución del diseño',2)
p('El monolito favorece transacciones locales, pocos saltos de red y una operación sencilla. La separación controller–service–repository permite optimizar consultas sin extraer servicios. Compartir base y recursos también crea riesgo de interferencia entre consultas y pedidos; estas campañas aisladas no cuantifican ese efecto.')
p('Si el GET continúa limitando la carga, evaluar un modelo de lectura materializado o una caché con invalidación de promociones y disponibilidad. Definir explícitamente cuánto dato obsoleto tolera el negocio. Separar lectura y escritura facilita escalarlas de forma diferente, pero exige sincronización y supervisión de retrasos.')
p('Una cola duradera para POST requiere contrato de estados, idempotencia y seguimiento hasta la persistencia. Responder 202 Accepted no equivale a un pedido creado y no cumple automáticamente ASR 2. La extracción a microservicios solo se justifica si el aislamiento y escalamiento independiente compensan el costo operativo y de consistencia; no elimina SQL ineficiente ni una base compartida saturada.')

page('11 Conclusiones')
p('El primer ASR que se incumple es el de latencia del GET. En la matriz principal ocurre desde alta carga Python; la comparación con 450 hilos JMeter no define un intervalo válido porque cambia el modelo. El refinamiento muestra cumplimiento repetido a 100 usuarios y resultados inestables a 150. Se debe tratar el límite como una zona de transición local y conservar margen operativo.')
p('POST tolera más concurrencia antes de superar su criterio de errores, pero ya acumula decenas de segundos de latencia antes del pico. En la matriz principal el cruce de 2 % está entre 7500 y 18000 usuarios; los resultados adicionales de la sección 7 precisan o matizan ese rango bajo un ramp-up común. El menor porcentaje de errores previo al pico no significa ausencia de degradación.')
p('La evidencia más directa de trabajo dominante está en PostgreSQL para GET: el plan recorre relaciones grandes y la rama histórica concentra gran parte del tiempo aislado. En POST, las 24 validaciones secuenciales y el pool de diez son candidatos claros; la saturación del generador en el pico impide asignar una causa única sin instrumentación adicional.')
p('Antes de un pico comercial, la primera mejora propuesta es validar productos por lote e introducir idempotencia en POST, porque ya se observaron pedidos persistidos sin confirmación al cliente. Los índices GET son una intervención acotada complementaria. Ambas mejoras deben compararse con la línea base; este informe no les atribuye ganancias que todavía no se midieron.')
p('El monolito permitió cumplir la operación normal medida y facilita cambios locales. Bajo carga, sus recursos compartidos no ofrecen aislamiento automático. La evidencia respalda optimizar consultas, administrar concurrencia y medir antes de migrar a microservicios. La arquitectura puede evolucionar hacia lecturas separadas cuando el costo y la consistencia estén justificados.')
p('Las conclusiones se limitan a un computador, generador local, tráfico por endpoint separado, tienda fija y datos sintéticos con relaciones parcialmente cíclicas. No equivalen a una certificación de producción ni a una prueba de la tasa abierta de 5000 pedidos por minuto. Una siguiente validación debe separar generador y servidor, alternar tratamientos y medir tráfico mixto con tasas objetivo explícitas.')

page('12 Prompts utilizados')
p('La generación del cliente Python ocurrió dentro de la tarea general con el agente. A continuación se conserva la instrucción completa que solicitó las pruebas y el entregable; no se atribuye a una sesión separada con Copilot.')
prompt_text=(LAB/'config/prompts.md').read_text()
sections=re.split(r'^## ',prompt_text,flags=re.M)
for sec in sections[1:]:
    lines=sec.strip().splitlines()
    if lines[0]=='Especificación aplicada por el agente al generar el script':page('12 Criterios de generación y revisión')
    h(lines[0],2)
    for line in lines[1:]:
        if line.strip():
            para=p(line)
            if lines[0]=='Instrucción recibida':
                for r in para.runs:r.font.size=Pt(10)

page('13 Anexos y evidencia de JMeter')
img('evidence/jmeter-config.png','Configuración del plan JMeter','Vista documental generada a partir del JMX: 450 hilos, ramp-up 50 s y un loop. No es una captura de la ventana Java.',6.8)
img('evidence/jmeter-get-percentiles.png','Estadísticas oficiales de JMeter para GET normal 1','El dashboard HTML generado desde el CSV muestra 450 muestras, p95 43 ms y p99 46 ms; el error es 0 %. Corresponde al agregado del reporte oficial.',6.8)

page('13 Evidencia de resultados POST')
img('evidence/jmeter-post-percentiles.png','Estadísticas oficiales de JMeter para POST normal 1','El reporte HTML muestra 450 muestras, p95 35 ms, p99 37.49 ms y 0 % de error. La tabla incluye el resumen y los percentiles del CSV.',6.8)
p('Los dashboards HTML oficiales, sus statistics.json y los planes JMX están en la entrega. Las vistas de configuración y Python se identifican por su origen. Los paneles Swing Summary Report y Aggregate Report no se presentan como capturados: la evidencia disponible corresponde al dashboard oficial generado desde las muestras reales.')
h('Validación de salud',2)
p('raw/health-headers.txt conserva HTTP 200 y raw/health.txt contiene Health Check OK. La captura original health.png se conserva como evidencia complementaria; no se amplía una página vacía para reemplazar el registro de la respuesta.')
h('Correspondencia con los archivos',2)
p('Cada nombre <endpoint>-<escenario>-<repetición> relaciona CSV, configuración, log y resumen. Las figuras de JMeter corresponden a get-normal-1 y post-normal-1. Los percentiles de las tablas fueron recalculados desde las muestras; la auditoría comprueba el número de corridas y la concordancia con los JSON Python.')

page('13 Evidencia del cliente Python')
img('evidence/python-get-alta.png','Comando y resultado archivado de GET alta 1','La vista de los registros guardados muestra 1500 usuarios, ramp-up 75 s, meseta 60 s, 37490 solicitudes y p99 de 9609.233 ms.',6.8)
img('evidence/python-post-pico.png','Comando y resultado archivado de POST pico 5','Se muestran 18000 usuarios, cuerpo grande, 116097 solicitudes, p95 y p99, throughput y 74.570 % de error. Es una vista del resumen guardado, no una nueva ejecución.',6.8)

page('13 Referencias y archivos de entrega')
h('Referencias técnicas',2)
refs=[
 'Guía Lab 2 Pruebas de carga al monolito de Cheapest y lab_2_warmup.md. Copias en reference/.',
 'Código inspeccionado: src/logistica/repositories/producto.repository.ts; src/logistica/services/pedido.service.ts; src/logistica/repositories/pedido.repository.ts; src/datasources/database.providers.ts; src/datasources/database-seeder.service.ts. Copias de referencia con extensión .txt en code-reference/.',
 'Evidencia SQL: raw/get-explain-analyze.txt, config/explain-get.sql y raw/*-baseline-indexes.csv. Límite local del pool: node_modules/pg-pool/index.js, valor por defecto 10.',
 '[1] GitHub Docs. Adding repository custom instructions for GitHub Copilot. https://docs.github.com/en/copilot/how-tos/copilot-on-github/customize-copilot/add-custom-instructions/add-repository-instructions',
 '[2] GitHub Docs. About protected branches. https://docs.github.com/en/repositories/configuring-branches-and-merges-in-your-repository/managing-protected-branches/about-protected-branches',
 '[3] GitHub Docs. About code owners. https://docs.github.com/en/repositories/managing-your-repositorys-settings-and-features/customizing-your-repository/about-code-owners']
for text in refs:
    para=p(text)
    for r in para.runs:r.font.size=Pt(10)
h('Mapa de la entrega',2)
table(['Carpeta o archivo','Contenido'],[
 ['Informe_Lab2_Corregido.docx','Informe completo con tablas, figuras y preguntas'],
 ['seeds y config','Variantes YAML, JMX, cuerpo POST, prompts y entorno'],
 ['scripts','Cliente, ejecución, auditoría y generación de artefactos'],
 ['raw','44 CSV originales, logs, JSON y snapshots'],
 ['refinement','Pruebas adicionales y sus configuraciones'],
 ['report','Consolidados GET y POST, resumen y auditoría'],
 ['graphs y evidence','Gráficas y capturas con sus fuentes'],
 ['reference y code-reference','Guías, checklist y código inspeccionado']], [2.4,4.6],9)
p('README_ENTREGA.md explica la reproducción. REVISION_CHECKLIST.md identifica las correcciones y la sustitución explícita de los paneles Java por reportes HTML oficiales. MANIFEST_SHA256.txt permite verificar los archivos del ZIP.')
d.core_properties.title='Pruebas de carga del monolito Cheapest'
d.core_properties.subject='Laboratorio 2 con auditoría de 44 corridas y refinamiento'
d.core_properties.author='José'
d.core_properties.keywords='Cheapest, ASR, carga, GET, POST, JMeter, Python'
out=OUT/'Informe_Lab2_Corregido.docx';d.save(out)
print(json.dumps({'output':str(out),'main_runs':len(rows),'refinement_runs':len(ref),'figures':figure}))
