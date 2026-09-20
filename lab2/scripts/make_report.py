from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from pathlib import Path
import csv,json
root=Path('lab2');d=Document();sec=d.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10)
for sty in ['Title','Heading 1','Heading 2']:
 styles[sty].font.name='Aptos Display';styles[sty].font.color.rgb=RGBColor(20,55,90)
def title(t):d.add_heading(t,0)
def h(t,l=1):d.add_heading(t,l)
def p(t):d.add_paragraph(t)
def table(headers,rows):
 t=d.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
 for c,x in zip(t.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=str(x)
 return t
def img(path,w=6.4):
 if Path(path).suffix.lower() in {'.png','.jpg','.jpeg'} and Path(path).exists():
  d.add_picture(str(path),width=Inches(w));d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

title('Laboratorio 2 Pruebas de carga al monolito Cheapest')
p('Informe experimental · Rama lab2-jose · 7 de septiembre de 2026')
p('Este informe documenta el diseño de datos, la preparación del entorno, las pruebas de carga de los endpoints GET y POST, la evidencia recolectada y el análisis arquitectónico. El resultado central es que el GET cumple con holgura en la operación normal medida por JMeter, pero supera ampliamente el p99 de 1000 ms cuando la concurrencia se eleva a 1500 usuarios y más. El POST se ejecutó con una base restaurada antes de cada corrida para evitar contaminación entre repeticiones.')
h('1. Entorno y reproducibilidad')
p('Repositorio local: Desktop/Uni/5thSemester/Arquisoft/cheapest-api. Rama: lab2-jose. Hardware: Apple M1, 8 CPU, 16 GB RAM, macOS 26.6.2, Node v26.8.1, Java 21. PostgreSQL 15 en Docker y NestJS local. La compilación de producción pasó correctamente. /health respondió Health Check OK.')
img(root/'evidence/health.png',5.8)
p('La descarga de Apache JMeter 5.6.3 se verificó con SHA-512. El plan se ejecutó en modo no gráfico y sus CSV se conservaron. El recurso load_test.jmx enlazado por la guía no estaba incluido localmente; por eso se reconstruyó un plan equivalente con los dos grupos y cuatro listeners solicitados.')
h('2. Diseño de datos')
p('Se definieron 75 tiendas y cinco zonas. La actividad se concentra en 20% de las tiendas y zonas, con peso de cabeza 80% para GET y 90% para POST. La variante GET usa 18.000 pedidos, 126.000 ítems y 22.500 asociaciones catálogo-producto para estresar la lectura histórica. La variante POST usa 4.500 pedidos, 31.500 ítems y 5.400 promociones; sus pedidos Python contienen 24 ítems. El seeder aplica Pareto a asignaciones de tiendas y catálogos, pero algunas relaciones siguen ciclos; esto queda documentado como limitación.')
table(['Entidad','GET ASR1','POST ASR2'],[['Tiendas','75','75'],['Zonas','5','5'],['Productos','1.800','1.800'],['Catálogos','125','125'],['Catálogo-producto','22.500','22.500'],['Promociones','3.600','5.400'],['Pedidos','18.000','4.500'],['Ítems de pedido','126.000','31.500'],['Distribución','Pareto 80/20','Pareto 90/20']])
img(root/'evidence/seed-config.html',5.8)
h('3. Matriz y resultados')
p('JMeter ejecutó 5, 30, 100 y 450 hilos con ramp-up de 5, 10, 20 y 50 s. Smoke, baja y media tuvieron una corrida; operación normal tuvo cinco. Python usa usuarios concurrentes, ramp-up y una meseta de 60 s. El throughput observado se reporta separado de los usuarios nominales.')
rows=[]
if (root/'report/results-summary.csv').exists():
 for r in csv.DictReader(open(root/'report/results-summary.csv')):
  rows.append([r['scenario'],r['endpoint'],r['rep'],r['users'],r['ramp_up'],r['duration'],r['p99_ms'],r['throughput_rps'],r['error_pct']])
table(['Escenario','Endpoint','Corrida','Usuarios','Ramp-up','Duración','p99 ms','req/s','Error %'],rows)
img(root/'evidence/jmeter-get-normal.png',6.0);img(root/'evidence/jmeter-get-percentiles.png',6.0)
img(root/'graphs/get-asr.png',6.0);img(root/'graphs/post-asr.png',6.0);img(root/'graphs/throughput.png',6.0)
h('4. Análisis y respuestas')
analysis=(root/'report/analysis-draft.md').read_text()
for line in analysis.splitlines():
 if line.startswith('# '):continue
 if line.startswith('## '):h(line[3:],1)
 elif line.startswith('[1]') or line.startswith('[2]') or line.startswith('[3]'):p(line)
 elif line.strip():p(line)
h('5. Evidencia y código entregado')
p('La carpeta lab2 contiene las semillas por escenario, el plan JMeter, el script load_test.py, el cuerpo grande de POST, los CSV y JSON de cada corrida, logs de JMeter y NestJS, telemetría de PostgreSQL, auditorías de conteos e índices, capturas, gráficas, prompts y guías de referencia. Los resultados no deben interpretarse como capacidad productiva: son mediciones locales, con un mock de tienda que espera 10 ms, sin cambios de índices ni pool.')
d.add_page_break();p('Fuentes externas para la Pregunta 4: GitHub Docs sobre instrucciones de repositorio para Copilot, ramas protegidas y CODEOWNERS. Las URL están en analysis-draft.md y prompts.md.')
out=root/'report/Informe_Lab2.docx';d.save(out)
print(out)
